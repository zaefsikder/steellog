"""Parse the two workout spreadsheets into a single normalized programs.json.

This is a BUILD-TIME script. It reads the source .xlsx files and emits
app/data/programs.json, which the FastAPI app loads at runtime. The runtime
does not depend on openpyxl.

Normalized shape:
    {
      "programs": [
        {
          "id": "hybrid",
          "name": "Hybrid Mobility + Hypertrophy",
          "subtitle": "3-4 Days / Week - 45-60 Min",
          "days": [
            {
              "id": "hybrid-day-a",
              "label": "Day A",
              "title": "Lower Body - Mobility + Squat Pattern + Hinge",
              "exercises": [
                {
                  "id": "hybrid-day-a-platz-stretch",
                  "name": "Platz Stretch",
                  "category": "Mobility",
                  "sets": "3",
                  "reps": "45-60s hold",
                  "weight": "",
                  "notes": "Deep squat hold, ..."
                }
              ]
            }
          ]
        }
      ]
    }

Run with:  uv run --group dev python scripts/parse_workouts.py
"""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "app" / "data" / "programs.json"


def slugify(*parts: str) -> str:
    text = " ".join(p for p in parts if p)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text


def clean(value) -> str:
    """Normalize a cell to a trimmed single-line string."""
    if value is None:
        return ""
    text = str(value).strip()
    # Collapse embedded newlines/whitespace runs.
    text = re.sub(r"\s+", " ", text)
    return text


def split_sets_reps(scheme: str) -> tuple[str, str]:
    """Split a "Sets x Reps" style string into (sets, reps).

    Examples:
        "4 x 5-8"            -> ("4", "5-8")
        "3 x 45-60s hold"    -> ("3", "45-60s hold")
        "4 x max (weighted)" -> ("4", "max (weighted)")
        "3-4 x 8-10 / leg"   -> ("3-4", "8-10 / leg")
    """
    if not scheme:
        return "", ""
    # Normalize the various multiplication marks to a plain 'x'.
    normalized = scheme.replace("×", "x").replace("X", "x")
    parts = normalized.split("x", 1)
    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip()
    return "", scheme.strip()


# ---------------------------------------------------------------------------
# hybrid_training.xlsx
# ---------------------------------------------------------------------------
# Sheet "Program Overview": a title row, then repeating blocks of
#   "  DAY A  -  Lower Body ..."  (day header, other columns blank)
#   "Exercise | Category | Sets x Reps | Notes / Cues"  (column header)
#   exercise rows...
def parse_hybrid(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["\U0001f4cb Program Overview"]

    rows = [[clean(c) for c in row] for row in ws.iter_rows(values_only=True)]

    subtitle = ""
    if rows and rows[0] and rows[0][0]:
        # e.g. "HYBRID MOBILITY + HYPERTROPHY  |  3-4 Days / Week  |  45-60 Min"
        title_cell = rows[0][0]
        segments = [s.strip() for s in title_cell.split("|") if s.strip()]
        subtitle = " - ".join(segments[1:]) if len(segments) > 1 else ""

    day_header_re = re.compile(r"^DAY\s+([A-Z])\s*[-–—]\s*(.*)$", re.IGNORECASE)

    days: list[dict] = []
    current: dict | None = None
    for row in rows:
        first = row[0]
        if not first:
            continue

        match = day_header_re.match(first)
        if match:
            letter = match.group(1).upper()
            remainder = match.group(2).strip(" -–—")
            current = {
                "id": f"hybrid-day-{letter.lower()}",
                "label": f"Day {letter}",
                "title": remainder,
                "exercises": [],
            }
            days.append(current)
            continue

        # Skip the per-block column header row.
        if first.lower() == "exercise":
            continue

        if current is None:
            continue

        name = first
        category = row[1] if len(row) > 1 else ""
        scheme = row[2] if len(row) > 2 else ""
        notes = row[3] if len(row) > 3 else ""
        sets, reps = split_sets_reps(scheme)
        current["exercises"].append(
            {
                "id": slugify(current["id"], name),
                "name": name,
                "category": category,
                "sets": sets,
                "reps": reps,
                "weight": "",
                "notes": notes,
            }
        )

    return {
        "id": "hybrid",
        "name": "Hybrid Mobility + Hypertrophy",
        "subtitle": subtitle,
        "days": days,
    }


# ---------------------------------------------------------------------------
# hypertrophy.xlsx
# ---------------------------------------------------------------------------
# Sheet "Workouts & Schedule": data lives in columns B..F (column A blank).
#   "Hypertrophy Training Program" (title)
#   "Upper"                        (section = day)
#   "Exercises | Sets | Reps | Wts | Notes"  (column header)
#   exercise rows...
def parse_hypertrophy(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Workouts & Schedule"]

    rows = [[clean(c) for c in row] for row in ws.iter_rows(values_only=True)]

    sections = {"upper", "lower", "push", "pull"}

    days: list[dict] = []
    current: dict | None = None
    for row in rows:
        # Data is offset one column; find the first non-empty label.
        label = ""
        label_col = 0
        for i, cell in enumerate(row):
            if cell:
                label = cell
                label_col = i
                break
        if not label:
            continue

        lower = label.lower()
        if lower == "hypertrophy training program":
            continue
        if lower == "exercises":
            continue

        if lower in sections:
            title = label.capitalize()
            current = {
                "id": f"hypertrophy-{lower}",
                "label": title,
                "title": f"{title} Day",
                "exercises": [],
            }
            days.append(current)
            continue

        if current is None:
            continue

        # Columns after the name: Sets, Reps, Wts, Notes.
        name = label
        rest = row[label_col + 1 :]
        sets = rest[0] if len(rest) > 0 else ""
        reps = rest[1] if len(rest) > 1 else ""
        weight = rest[2] if len(rest) > 2 else ""
        notes = rest[3] if len(rest) > 3 else ""
        current["exercises"].append(
            {
                "id": slugify(current["id"], name),
                "name": name,
                "category": "",
                "sets": sets,
                "reps": reps,
                "weight": weight,
                "notes": notes,
            }
        )

    return {
        "id": "hypertrophy",
        "name": "Hypertrophy",
        "subtitle": "Upper / Lower / Push / Pull",
        "days": days,
    }


def main() -> None:
    programs = [
        parse_hypertrophy(ROOT / "hypertrophy.xlsx"),
        parse_hybrid(ROOT / "hybrid_training.xlsx"),
    ]
    payload = {"programs": programs}

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n")

    # Summary for a quick sanity check.
    for program in programs:
        total = sum(len(d["exercises"]) for d in program["days"])
        print(
            f"{program['id']:12s} {len(program['days'])} days, "
            f"{total} exercises -> {program['name']}"
        )
    print(f"\nWrote {OUT_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
